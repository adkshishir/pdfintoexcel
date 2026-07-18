"""Bypass broken Unicode text layers (common for some Devanagari PDFs)."""

from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from app.pipeline import digital_words as digital_words_module
from app.pipeline import orchestrator
from app.pipeline.types import WordBox


def test_trust_pdf_text_false_routes_digital_through_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = tmp_path / "digital.pdf"
    doc = fitz.open()
    p = doc.new_page(width=400, height=400)
    p.insert_text((40, 40), "Item")
    p.insert_text((180, 40), "Quantity")
    for i, name in enumerate(("Apple", "Banana", "Cherry")):
        p.insert_text((40, 70 + 30 * i), name)
        p.insert_text((180, 70 + 30 * i), str(i + 1))
    doc.save(pdf)
    doc.close()

    calls: list[str] = []

    def spy_extract_digital_words(*args, **kwargs):  # noqa: ANN002,ANN003
        calls.append("digital_words")
        return digital_words_module.extract_digital_words(*args, **kwargs)

    monkeypatch.setattr(orchestrator, "extract_digital_words", spy_extract_digital_words)

    ocr_calls = {"n": 0}

    def fake_extract_ocr(path, *, mode, page_indices=None, recognizer=None, **kwargs):  # noqa: ANN001, ANN003
        ocr_calls["n"] += 1
        boxes = [
            WordBox("Item", 40, 40, 80, 10, page=0, confidence=0.95),
            WordBox("Quantity", 180, 40, 90, 10, page=0, confidence=0.95),
            WordBox("Apple", 40, 70, 60, 10, page=0, confidence=0.93),
            WordBox("3", 180, 70, 20, 10, page=0, confidence=0.93),
            WordBox("Banana", 40, 100, 70, 10, page=0, confidence=0.93),
            WordBox("5", 185, 100, 15, 10, page=0, confidence=0.93),
            WordBox("Cherry", 40, 130, 65, 10, page=0, confidence=0.93),
            WordBox("7", 185, 130, 15, 10, page=0, confidence=0.93),
        ]
        return (boxes, 1)

    monkeypatch.setattr(orchestrator, "extract_ocr", fake_extract_ocr)

    out = tmp_path / "out.xlsx"
    orchestrator.run_pipeline(pdf, out, mode="fast", trust_pdf_text=False)
    assert ocr_calls["n"] == 1
    assert calls == []
    assert out.exists()


def test_nep_reference_pdf_prefers_tesseract_over_embedded_when_forced(tmp_path: Path) -> None:
    """Regression for PDFs whose text layer maps Devanagari to wrong Unicode.

    Requires ``NEPALI_OCR_TEST_PDF`` (path to suntalajaat_book_2076.pdf or similar),
    system ``tesseract`` with Nepali traineddata, and OCR_ENGINE=tesseract while
    the test runs.
    """
    from app.config import get_settings

    pdf_src = os.environ.get("NEPALI_OCR_TEST_PDF")
    if not pdf_src:
        pytest.skip("Set NEPALI_OCR_TEST_PDF=/path/to/suntalajaat_book_2076.pdf")
    src = Path(pdf_src).expanduser()
    if not src.is_file():
        pytest.skip(f"NEPALI_OCR_TEST_PDF missing: {src}")

    if not shutil.which("tesseract"):
        pytest.skip("tesseract not on PATH")

    langs = subprocess.run(
        ["tesseract", "--list-langs"],
        capture_output=True,
        check=False,
        text=True,
    )
    if langs.returncode != 0 or "nep" not in langs.stdout.lower():
        pytest.skip("tesseract Nepali lang pack not installed (tesseract-ocr-nep)")

    old_engine = os.environ.get("OCR_ENGINE")
    try:
        os.environ["OCR_ENGINE"] = "tesseract"
        get_settings.cache_clear()

        one_page = tmp_path / "np-p0.pdf"
        doc = fitz.open(src)
        d = fitz.open()
        d.insert_pdf(doc, from_page=0, to_page=0)
        d.save(one_page)
        d.close()
        doc.close()

        def flatten(path: Path) -> str:
            wb = load_workbook(path, read_only=True, data_only=True)
            parts: list[str] = []
            for name in wb.sheetnames:
                if name.startswith("_conf_"):
                    continue
                ws = wb[name]
                for row in ws.iter_rows(values_only=True):
                    parts.extend("" if v is None else str(v) for v in row)
            wb.close()
            return "\n".join(parts)

        nep_ok = "\u0928\u0947\u092a\u093e\u0932"  # नेपाल
        corrupt_city = "\u0928\u0947\u0929\u093e\u0930"  # नेऩार (bad cmap)

        out_embed = tmp_path / "embed.xlsx"
        orchestrator.run_pipeline(one_page, out_embed, mode="accurate")

        out_ocr = tmp_path / "ocr.xlsx"
        orchestrator.run_pipeline(
            one_page,
            out_ocr,
            mode="accurate",
            trust_pdf_text=False,
        )

        te = flatten(out_embed)
        to = flatten(out_ocr)

        if nep_ok not in te and nep_ok not in to and corrupt_city not in te:
            pytest.skip("Reference PDF typography changed — extend heuristics for this fixture")

        corrupt_down = te.count(corrupt_city) > to.count(corrupt_city)
        nep_up = to.count(nep_ok) > te.count(nep_ok)
        assert corrupt_down or nep_up, (
            f"OCR regression: corrupt={te.count(corrupt_city)}->{to.count(corrupt_city)} "
            f"nep_ok={te.count(nep_ok)}->{to.count(nep_ok)}"
        )
    finally:
        if old_engine is None:
            os.environ.pop("OCR_ENGINE", None)
        else:
            os.environ["OCR_ENGINE"] = old_engine
        get_settings.cache_clear()
