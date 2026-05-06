"""Regression tests for full_document structured classification / export."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.pipeline.content_classifier import (
    _filter_grid_banner_header_lines,
    _strip_redundant_grid_header_body_blocks,
    classify_document,
)
from app.pipeline.structured_exporter import export_structured_document
from app.pipeline.types import CleanTable, ContentBlock, DocumentContent, RawCell, RawTable, WordBox


def _minimal_raw_table(
    page: int,
    y0: float,
    *,
    x0: float = 50.0,
    x1: float = 400.0,
    y1: float = 360.0,
) -> RawTable:
    cells = [
        RawCell(0, 0, "Col A", row_span=1, col_span=1),
        RawCell(0, 1, "Col B", row_span=1, col_span=1),
        RawCell(1, 0, "a1", row_span=1, col_span=1),
        RawCell(1, 1, "b1", row_span=1, col_span=1),
        RawCell(2, 0, "a2", row_span=1, col_span=1),
        RawCell(2, 1, "b2", row_span=1, col_span=1),
        RawCell(3, 0, "a3", row_span=1, col_span=1),
        RawCell(3, 1, "b3", row_span=1, col_span=1),
    ]
    return RawTable(
        page=page,
        bbox=(x0, y0, x1, y1),
        cells=cells,
        n_rows=4,
        n_cols=2,
        row_y_centers=[y0 + 12, y0 + 36, y0 + 60, y0 + 84],
    )


def _clean_from_raw(rt: RawTable) -> CleanTable:
    return CleanTable(
        sheet_name="T",
        headers=["Col A", "Col B"],
        rows=[["a1", "b1"], ["a2", "b2"], ["a3", "b3"]],
        confidence=[[1.0, 1.0], [1.0, 1.0], [1.0, 1.0]],
        column_types=["text", "text"],
        source_page=rt.page,
        source_y=rt.bbox[1],
    )


def test_canonical_header_first_page_with_banner_not_hardcoded_page_zero() -> None:
    """Running header must not disappear when page 0 is a cover without that banner."""
    # Three pages so the repeated line clears the 55% + min-2-pages gate.
    h = 12.0
    boxes = [
        WordBox("Cover", 72, 50, 120, h, page=0),
        WordBox("Acme", 72, 50, 80, h, page=1),
        WordBox("Acme", 72, 50, 80, h, page=2),
    ]
    raw = [_minimal_raw_table(1, 200.0), _minimal_raw_table(2, 200.0)]
    clean = [_clean_from_raw(raw[0]), _clean_from_raw(raw[1])]

    doc = classify_document(boxes, raw, clean, page_count=3, pdf_path=None)
    assert doc.header_lines == ["Acme"]


def test_pre_table_block_rendered_like_footer_rows(tmp_path: Path) -> None:
    """Intro lines above the first table use per-line merged rows (footer-style)."""
    h = 12.0
    boxes = [
        WordBox("Acme", 72, 40, 80, h, page=0),
        WordBox("Acme", 72, 40, 80, h, page=1),
        WordBox("LineOne", 72, 120, 200, h, page=1),
        WordBox("LineTwo", 72, 140, 220, h, page=1),
    ]
    raw = [_minimal_raw_table(1, 220.0)]
    clean = [_clean_from_raw(raw[0])]
    doc = classify_document(boxes, raw, clean, page_count=2, pdf_path=None)

    pre = [b for b in doc.body_blocks if b.kind == "pre_table"]
    assert len(pre) == 1
    assert pre[0].lines == ["LineOne", "LineTwo"]

    out = tmp_path / "doc.xlsx"
    export_structured_document(doc, out)
    wb = load_workbook(out)
    ws = wb.active
    flat: list[str | int | float | None] = []
    for row in ws.iter_rows(values_only=True):
        flat.extend(v for v in row if v is not None and str(v).strip() != "")

    idx_one = flat.index("LineOne")
    idx_two = flat.index("LineTwo")
    idx_hdr = flat.index("Col A")
    assert idx_one < idx_two < idx_hdr


def test_preamble_kept_when_summary_table_bbox_overlays_header_area() -> None:
    """Upper summary grid must not consume Monzo-style statement text above the tx grid."""
    h = 11.0
    boxes = [
        WordBox("Business", 72, 55, 130, h, page=0),
        WordBox("Account", 150, 55, 80, h, page=0),
        WordBox("GridStart", 72, 270, 180, h, page=0),
    ]
    raw_top = RawTable(
        page=0,
        bbox=(40.0, 45.0, 400.0, 100.0),
        cells=[
            RawCell(0, 0, "S1", row_span=1, col_span=1),
            RawCell(0, 1, "S2", row_span=1, col_span=1),
            RawCell(1, 0, "x", row_span=1, col_span=1),
            RawCell(1, 1, "y", row_span=1, col_span=1),
            RawCell(2, 0, "x", row_span=1, col_span=1),
            RawCell(2, 1, "y", row_span=1, col_span=1),
            RawCell(3, 0, "x", row_span=1, col_span=1),
            RawCell(3, 1, "y", row_span=1, col_span=1),
        ],
        n_rows=4,
        n_cols=2,
        row_y_centers=[52.0, 64.0, 76.0, 88.0],
    )
    raw_main = _minimal_raw_table(0, 250.0, y1=400.0)
    clean = [_clean_from_raw(raw_main)]
    doc = classify_document(boxes, [raw_top, raw_main], clean, page_count=1, pdf_path=None)
    joined = " ".join(
        " ".join(b.lines) for b in doc.body_blocks if b.kind != "table"
    )
    assert "Business" in joined and "Account" in joined
    assert "GridStart" not in joined


def test_single_table_cut_at_date_description_header_row() -> None:
    """One holistic RawTable: narrative lives above the Date/Description header row."""
    h = 10.0
    boxes = [
        WordBox("Statement", 50, 52, 120, h, page=0),
        WordBox("Date", 50, 278, 44, h, page=0),
        WordBox("Desc", 120, 278, 200, h, page=0),
        WordBox("Amount", 400, 278, 70, h, page=0),
        WordBox("Bal", 500, 278, 50, h, page=0),
        WordBox("R1", 50, 310, 40, h, page=0),
    ]
    y0, y1 = 40.0, 380.0
    cells = [
        RawCell(0, 0, "Statement title row", row_span=1, col_span=1),
        RawCell(1, 0, "Secondary", row_span=1, col_span=1),
        RawCell(2, 0, "Date", row_span=1, col_span=1),
        RawCell(2, 1, "Description (GBP)", row_span=1, col_span=1),
        RawCell(2, 2, "Amount (GBP)", row_span=1, col_span=1),
        RawCell(2, 3, "Balance", row_span=1, col_span=1),
        RawCell(3, 0, "28/07/2025", row_span=1, col_span=1),
        RawCell(3, 1, "Payee", row_span=1, col_span=1),
        RawCell(3, 2, "1.00", row_span=1, col_span=1),
        RawCell(3, 3, "100.00", row_span=1, col_span=1),
    ]
    raw = RawTable(
        page=0,
        bbox=(y0, y0, 560.0, y1),
        cells=cells,
        n_rows=4,
        n_cols=4,
        row_y_centers=[60.0, 130.0, 283.0, 315.0],
    )
    clean = [CleanTable(
        sheet_name="Stmt",
        headers=["Date", "Description (GBP)", "Amount (GBP)", "Balance"],
        rows=[["28/07/2025", "Payee", "1.00", "100.00"]],
        confidence=[[1.0, 1.0, 1.0, 1.0]],
        column_types=["text", "text", "text", "text"],
        source_page=0,
        source_y=raw.bbox[1],
    )]
    doc = classify_document(boxes, [raw], clean, page_count=1, pdf_path=None)
    narrative = " ".join(" ".join(b.lines) for b in doc.body_blocks if b.kind != "table")
    assert "Statement" in narrative
    assert "Date" not in narrative
    assert "R1" not in narrative


def test_document_header_omits_repeated_transaction_column_titles() -> None:
    """Grid column header repeated in PDF header-zone must not become DOCUMENT HEADER."""
    ct = CleanTable(
        sheet_name="Tx",
        headers=["Date", "Description (GBP)", "Amount (GBP)", "Balance"],
        rows=[["28/07/2025", "x", "1", "2"]],
        confidence=[[1.0, 1.0, 1.0, 1.0]],
        column_types=["text", "text", "text", "text"],
        source_page=0,
        source_y=200.0,
    )
    lines = [
        "Date Description (GBP) Amount (GBP) Balance",
        "Monzo Bank Ltd",
    ]
    assert _filter_grid_banner_header_lines(lines, [ct]) == ["Monzo Bank Ltd"]


def test_body_drops_standalone_column_header_paragraph() -> None:
    ct = CleanTable(
        sheet_name="Tx",
        headers=["Date", "Description (GBP)", "Amount (GBP)", "Balance"],
        rows=[["28/07/2025", "x", "1", "2"]],
        confidence=[[1.0, 1.0, 1.0, 1.0]],
        column_types=["text", "text", "text", "text"],
        source_page=1,
        source_y=250.0,
    )
    blocks = [
        ContentBlock(
            kind="pre_table",
            page=1,
            y_start=40.0,
            lines=["Date Description (GBP) Amount (GBP) Balance"],
            table=None,
        ),
        ContentBlock(
            kind="paragraph",
            page=1,
            y_start=80.0,
            lines=["Real preamble"],
            table=None,
        ),
    ]
    out = _strip_redundant_grid_header_body_blocks(blocks, [ct])
    assert len(out) == 1
    assert out[0].lines == ["Real preamble"]


def test_export_per_page_writes_one_sheet_per_pdf_page(tmp_path: Path) -> None:
    from app.pipeline.structured_exporter import export_structured_document_per_page

    doc = DocumentContent(
        header_lines=[],
        footer_lines=[],
        body_blocks=[
            ContentBlock(kind="paragraph", page=0, y_start=10, lines=["Alpha"], table=None),
            ContentBlock(kind="paragraph", page=1, y_start=10, lines=["Beta"], table=None),
        ],
        page_count=2,
    )
    out = tmp_path / "multi.xlsx"
    export_structured_document_per_page(doc, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Page_1", "Page_2"]
