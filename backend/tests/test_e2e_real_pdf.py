"""End-to-end test against the real MonzoBus.pdf fixture.

Goes all the way: HTTP upload → Celery (eager) → reconstruction → cleaning →
Excel export → HTTP download → assertions on actual cell values.

Skipped if the fixture isn't checked in.
"""

from __future__ import annotations

import io
import os
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
MONZO_PDF = REPO_ROOT / "MonzoBus.pdf"


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    storage_root = tmp_path_factory.mktemp("storage")
    db_path = tmp_path_factory.mktemp("db") / "e2e.db"
    os.environ["DATABASE_URL"] = f"sqlite:///{db_path}"
    os.environ["STORAGE_BACKEND"] = "local"
    os.environ["STORAGE_LOCAL_ROOT"] = str(storage_root)
    os.environ["REDIS_URL"] = "memory://"

    from app import config, storage
    from app.models import database
    config.get_settings.cache_clear()
    storage.get_storage.cache_clear()
    database._engine = None
    database._SessionLocal = None

    from app.models.database import Base, get_engine
    from app.models import job as _job  # noqa: F401
    Base.metadata.create_all(get_engine())

    from app.queue.celery_app import celery_app
    celery_app.conf.broker_url = "memory://"
    celery_app.conf.result_backend = "cache+memory://"
    celery_app.conf.task_always_eager = True

    from app.api import limiter as lim
    lim.enabled = False

    from app.main import app
    return TestClient(app)


@pytest.mark.skipif(not MONZO_PDF.exists(), reason="MonzoBus.pdf fixture not present")
def test_e2e_monzo_accurate_mode(client):
    """Upload MonzoBus.pdf in accurate mode → download → assert known transactions."""
    with MONZO_PDF.open("rb") as fh:
        r = client.post(
            "/api/jobs",
            files={"file": ("MonzoBus.pdf", fh, "application/pdf")},
            data={"mode": "accurate"},
        )
    assert r.status_code == 202, r.text
    job_id = r.json()["id"]

    final = client.get(f"/api/jobs/{job_id}").json()
    assert final["status"] == "completed", final
    assert final["pdf_type"] == "digital"
    assert final["page_count"] == 24
    assert final["metrics"]["table_count"] == 1   # 23 pages collapsed to 1

    dl = client.get(f"/api/jobs/{job_id}/download")
    assert dl.status_code == 200
    out_bytes = dl.content
    assert out_bytes[:2] == b"PK"   # xlsx is a zip

    wb = load_workbook(io.BytesIO(out_bytes))
    data_sheets = [n for n in wb.sheetnames if not n.startswith("_conf_")]
    conf_sheets = [n for n in wb.sheetnames if n.startswith("_conf_")]
    assert len(data_sheets) == 1, data_sheets
    assert len(conf_sheets) == 1, conf_sheets

    ws = wb[data_sheets[0]]

    # Locate the table header row (preamble rows may precede it).
    table_header_row = None
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        if vals[0] == "Date" and "Description" in vals:
            table_header_row = row[0].row
            break
    assert table_header_row is not None, "Could not find table header row in sheet"
    headers = [c.value for c in ws[table_header_row]]
    assert headers == ["Date", "Description", "(GBP) Amount", "(GBP) Balance"]

    body = [[c.value for c in r] for r in ws.iter_rows(min_row=table_header_row + 1)]
    assert len(body) > 100, "expected many transaction rows"

    # Type assertions — accurate mode should preserve column types.
    typed_dates    = sum(1 for r in body if isinstance(r[0], date))
    typed_amounts  = sum(1 for r in body if isinstance(r[2], (int, float)))
    typed_balances = sum(1 for r in body if isinstance(r[3], (int, float)))
    assert typed_dates    > 100, f"only {typed_dates} typed dates"
    assert typed_amounts  > 100, f"only {typed_amounts} typed amounts"
    assert typed_balances > 100, f"only {typed_balances} typed balances"

    # Spot-check transactions that should be present after Phase 5's
    # cleanup + Phase 6's continuation merging. NOTE: page 0 of the PDF
    # carries the cover-page summary AND ~8 most-recent transactions in the
    # same page; `_looks_like_table` drops the whole page as cover-noise,
    # so transactions on page 0 (e.g. 25/07 WICKES, 26/07 SumUp) are
    # currently lost. Two-table-per-page handling would recover them —
    # see docs/phase-10-tests-and-benchmark.md "known-loss".
    descriptions = " | ".join(str(r[1]) for r in body if r[1])
    assert "Amazon" in descriptions or "AMAZON" in descriptions
    assert "TFL" in descriptions or "TfL" in descriptions
    assert "SAINSBURYS" in descriptions or "Sainsburys" in descriptions

    # Confidence sheet covers the table data rows.
    # body may contain extra appended footer rows — so conf_rows ≤ len(body).
    cws = wb[conf_sheets[0]]
    conf_rows = sum(1 for _ in cws.iter_rows(min_row=2))
    assert conf_rows <= len(body), f"conf rows {conf_rows} > body rows {len(body)}"
    assert conf_rows > 100, f"too few conf rows: {conf_rows}"
