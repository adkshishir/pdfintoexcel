"""Figures sheet embedding for ``image_export=figures``."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.pipeline.exporter import export_excel
from app.pipeline.figures_sheet import append_figures_sheet_from_pdf
from app.pipeline.types import CleanTable


def _pdf_with_embedded_png(tmp_path: Path) -> Path:
    import fitz
    from PIL import Image

    png = tmp_path / "blob.png"
    Image.new("RGB", (80, 50), color=(10, 120, 200)).save(png)
    pdf = tmp_path / "one_img.pdf"
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    r = fitz.Rect(50, 50, 200, 200)
    page.insert_image(r, filename=str(png))
    doc.save(pdf)
    doc.close()
    return pdf


def test_append_figures_sheet_adds_worksheet(tmp_path: Path) -> None:
    pdf = _pdf_with_embedded_png(tmp_path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Data")
    ws["A1"] = "x"
    stats = append_figures_sheet_from_pdf(wb, pdf)
    assert stats.get("embedded", 0) >= 1
    assert any(n.startswith("Figures") for n in wb.sheetnames)


def test_export_excel_empty_tables_with_figures(tmp_path: Path) -> None:
    pdf = _pdf_with_embedded_png(tmp_path)
    out = tmp_path / "out.xlsx"
    em: dict = {}
    export_excel(
        [],
        out,
        image_export="figures",
        pdf_path=pdf,
        export_metrics=em,
    )
    wb = openpyxl.load_workbook(out)
    assert "empty" in wb.sheetnames
    assert any(n.startswith("Figures") for n in wb.sheetnames)
    assert em.get("figures_embedded", 0) >= 1


def test_export_excel_with_table_and_figures(tmp_path: Path) -> None:
    pdf = _pdf_with_embedded_png(tmp_path)
    ct = CleanTable(
        sheet_name="T1",
        headers=["A", "B"],
        rows=[["1", "2"]],
        confidence=[[1.0, 1.0]],
        column_types=["text", "text"],
        row_gaps=[float("inf")],
        source_page=0,
        source_y=0.0,
    )
    out = tmp_path / "tbl.xlsx"
    export_excel(
        [ct],
        out,
        image_export="figures",
        pdf_path=pdf,
        export_metrics={},
    )
    wb = openpyxl.load_workbook(out)
    assert "T1" in wb.sheetnames
    assert any(n.startswith("Figures") for n in wb.sheetnames)
