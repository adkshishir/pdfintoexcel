"""``image_export=only`` short-circuits the table/OCR pipeline."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.pipeline import orchestrator


def _pdf_with_embedded_png(tmp_path: Path) -> Path:
    import fitz
    from PIL import Image

    png = tmp_path / "blob.png"
    Image.new("RGB", (80, 50), color=(10, 120, 200)).save(png)
    pdf = tmp_path / "one_img.pdf"
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    r = fitz.Rect(50, 50, 200, 200)
    page.insert_image(r, filename=str(png))
    doc.save(pdf)
    doc.close()
    return pdf


def test_image_export_only_skips_table_pipeline(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = _pdf_with_embedded_png(tmp_path)
    out = tmp_path / "only.xlsx"

    def fail_export(*args, **kwargs):  # noqa: ANN002,ANN003
        raise AssertionError("export_excel should not run for image_export=only")

    monkeypatch.setattr(orchestrator, "export_excel", fail_export)

    result = orchestrator.run_pipeline(pdf, out, image_export="only")
    assert out.exists()
    assert result.tables == []
    wb = openpyxl.load_workbook(out)
    assert any(n.startswith("Figures") for n in wb.sheetnames)
    assert result.export_metrics.get("figures_embedded", 0) >= 1


def test_image_export_only_empty_pdf(tmp_path: Path) -> None:
    import fitz

    pdf = tmp_path / "blank.pdf"
    doc = fitz.open()
    doc.new_page(width=200, height=200)
    doc.save(pdf)
    doc.close()

    out = tmp_path / "empty_only.xlsx"
    result = orchestrator.run_pipeline(pdf, out, image_export="only")
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    assert "empty" in wb.sheetnames or any(n.startswith("Figures") for n in wb.sheetnames)
    assert result.export_metrics.get("figures_embedded", 0) == 0
